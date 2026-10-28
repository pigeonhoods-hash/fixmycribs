"""Backend tests for FixMyCribs: property capacity/occupancy + bills export + AI import."""
import io
import os
import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE_URL:
    # fallback read from frontend/.env
    with open("/app/frontend/.env") as f:
        for ln in f:
            if ln.startswith("REACT_APP_BACKEND_URL="):
                BASE_URL = ln.split("=", 1)[1].strip().rstrip("/")

API = f"{BASE_URL}/api"
EMAIL = "landlord@test.com"
PASSWORD = "Test@1234"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{API}/auth/login", json={"email": EMAIL, "password": PASSWORD}, timeout=30)
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text}"
    data = r.json()
    tok = data.get("access_token") or data.get("token")
    assert tok, f"no token in {data}"
    return tok


@pytest.fixture(scope="module")
def auth(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def properties(auth):
    r = requests.get(f"{API}/properties", headers=auth, timeout=30)
    assert r.status_code == 200
    props = r.json()
    assert isinstance(props, list) and len(props) > 0
    return props


# ---------- Property capacity / occupancy ----------
class TestPropertiesCapacity:
    def test_list_properties_has_capacity_and_occupancy(self, properties):
        p = properties[0]
        assert "capacity" in p
        assert "capacity_type" in p
        assert "occupancy_pct" in p  # may be None if capacity 0
        assert "active_tenant_count" in p

    def test_lekki_court_capacity_and_occupancy(self, properties):
        lekki = next((p for p in properties if p.get("name") == "Lekki Court"), None)
        assert lekki, "Lekki Court not found"
        assert lekki["capacity"] > 0
        assert lekki["capacity_type"] in ("occupants", "beds", "flats", "rooms", "units")
        assert lekki["occupancy_pct"] is not None
        expected = round(min(100.0, (lekki["active_tenant_count"] / lekki["capacity"]) * 100), 1)
        assert abs(lekki["occupancy_pct"] - expected) < 0.01

    def test_create_property_with_capacity_and_verify(self, auth):
        payload = {
            "name": "TEST_CapProp",
            "address": "1 Test Rd",
            "type": "residential",
            "capacity": 8,
            "capacity_type": "rooms",
        }
        r = requests.post(f"{API}/properties", headers=auth, json=payload, timeout=30)
        assert r.status_code in (200, 201), r.text
        created = r.json()
        pid = created["id"]
        assert created["capacity"] == 8
        assert created["capacity_type"] == "rooms"

        # GET to verify persistence
        g = requests.get(f"{API}/properties/{pid}", headers=auth, timeout=30)
        assert g.status_code == 200
        gp = g.json()
        assert gp["capacity"] == 8
        assert gp["capacity_type"] == "rooms"
        # occupancy should be 0 (no tenants)
        assert gp.get("occupancy_pct") == 0 or gp.get("occupancy_pct") == 0.0

        # cleanup
        requests.delete(f"{API}/properties/{pid}", headers=auth, timeout=30)


# ---------- Bills Export ----------
class TestBillsExport:
    def test_export_all_periods_all_properties(self, auth):
        r = requests.get(f"{API}/bills/export", headers=auth, params={"period": "all"}, timeout=60)
        assert r.status_code == 200, r.text
        assert "spreadsheet" in r.headers.get("content-type", "")
        wb = load_workbook(io.BytesIO(r.content))
        assert len(wb.sheetnames) >= 1
        # Content-Disposition filename
        assert "bills" in r.headers.get("content-disposition", "").lower()

    def test_export_specific_month(self, auth):
        r = requests.get(f"{API}/bills/export", headers=auth,
                         params={"period": "month", "month": "2026-06"}, timeout=60)
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content))
        # Either has one sheet named "Jun 2026" or "No data"
        assert len(wb.sheetnames) >= 1

    def test_export_specific_year(self, auth):
        r = requests.get(f"{API}/bills/export", headers=auth,
                         params={"period": "year", "year": "2026"}, timeout=60)
        assert r.status_code == 200

    def test_export_single_property_and_tenant(self, auth, properties):
        lekki = next((p for p in properties if p.get("name") == "Lekki Court"), None)
        assert lekki
        # tenants
        tr = requests.get(f"{API}/tenants", headers=auth, timeout=30)
        assert tr.status_code == 200
        tenants = [t for t in tr.json() if t.get("property_id") == lekki["id"]]
        assert tenants
        tid = tenants[0]["id"]
        r = requests.get(f"{API}/bills/export", headers=auth,
                         params={"period": "all", "property_id": lekki["id"], "tenant_id": tid},
                         timeout=60)
        assert r.status_code == 200
        assert "spreadsheet" in r.headers.get("content-type", "")


# ---------- AI Import ----------
def _make_import_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Bills"
    ws.append(["Tenant Name", "Apartment", "Bill Type", "Amount", "Amount Paid",
               "Units", "Rate", "Date"])
    ws.append(["John Doe", "A1", "electricity", 12000, 12000, 200, 60, "2026-06-15"])
    ws.append(["TEST_Ghost Tenant", "Z9", "service", 5000, 0, 0, 0, "2026-06-20"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestBillImport:
    def test_analyze_and_commit(self, auth, properties):
        lekki = next((p for p in properties if p.get("name") == "Lekki Court"), None)
        assert lekki
        xlsx = _make_import_xlsx()
        files = {"file": ("bills.xlsx", xlsx,
                          "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        data = {"property_id": lekki["id"]}
        r = requests.post(f"{API}/bills/import/analyze", headers=auth,
                          files=files, data=data, timeout=90)
        assert r.status_code == 200, r.text
        preview = r.json()
        assert preview["property_id"] == lekki["id"]
        assert preview["summary"]["total_rows"] >= 2
        assert preview["summary"]["matched_count"] >= 1
        assert preview["summary"]["unmatched_count"] >= 1

        # Commit: import matched, skip unmatched (avoid polluting DB)
        rows = []
        for r_ in preview.get("matched", []):
            rows.append({**r_, "action": "import", "create_tenant": False})
        for r_ in preview.get("unmatched", []):
            rows.append({**r_, "action": "skip", "create_tenant": False})
        c = requests.post(f"{API}/bills/import/commit", headers=auth,
                         json={"property_id": lekki["id"], "rows": rows}, timeout=60)
        assert c.status_code == 200, c.text
        result = c.json()
        assert result["created_bills"] >= 1
        assert result["skipped"] >= 1
